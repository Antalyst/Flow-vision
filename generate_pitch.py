import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Colors
DEEP_SLATE = RGBColor(0x0F, 0x17, 0x2A)
ACCENT_CYAN = RGBColor(0x06, 0xB6, 0xD4)
ACCENT_EMERALD = RGBColor(0x10, 0xB9, 0x81)
CRISP_WHITE = RGBColor(0xF8, 0xFA, 0xFC)

def apply_dark_theme(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DEEP_SLATE

def add_title_slide(prs):
    slide_layout = prs.slide_layouts[6] # blank layout
    slide = prs.slides.add_slide(slide_layout)
    apply_dark_theme(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "FlowVision"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    
    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.2), Inches(8), Inches(1))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Decentralized Intelligence. Unified Control."
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(28)
    p2.font.color.rgb = CRISP_WHITE
    
    # Tagline
    txBox3 = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(1))
    tf3 = txBox3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "AI-Powered Document Tracking & Predictive Mesh for Municipal Enterprises"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(20)
    p3.font.color.rgb = ACCENT_EMERALD
    
    # Footer
    txBox4 = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.5))
    tf4 = txBox4.text_frame
    p4 = tf4.paragraphs[0]
    p4.text = "Team Immortals | DASIG.AI Pitching Competition 2026 (SMX Convention Center)"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(14)
    p4.font.color.rgb = CRISP_WHITE

def add_content_slide(prs, title, points, footer=""):
    slide_layout = prs.slide_layouts[6] # blank layout
    slide = prs.slides.add_slide(slide_layout)
    apply_dark_theme(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    
    # Content
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9), Inches(4.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    
    for i, point in enumerate(points):
        p = tf2.add_paragraph() if i > 0 else tf2.paragraphs[0]
        p.text = "• " + point
        p.font.size = Pt(22)
        p.font.color.rgb = CRISP_WHITE
        p.space_after = Pt(14)
        if "Key Takeaway" in point or "Core Summary" in point:
            p.font.color.rgb = ACCENT_EMERALD
            p.font.bold = True
            
    if footer:
        txBox3 = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.5))
        tf3 = txBox3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = footer
        p3.alignment = PP_ALIGN.CENTER
        p3.font.size = Pt(14)
        p3.font.color.rgb = CRISP_WHITE
            
prs = Presentation()

# Slide 1
add_title_slide(prs)

# Slide 2
add_content_slide(prs, "Zero-Visibility Tracking in Local Governance", [
    "Administrative Bottlenecks: Offices are overwhelmed by unexpected document surges without early warning systems.",
    "Fragmented Handoffs: Physical movement relies on manual logbooks, causing lost files and zero real-time visibility.",
    "Zero SLA Accountability: Impossible to identify delays in the routing mesh when deadlines are missed.",
    "Key Takeaway: Traditional systems only record past delays—they cannot prevent them."
])

# Slide 3
add_content_slide(prs, "Moving from Passive Logs to Active Intelligence", [
    "Beyond Digital CRUD: Conventional systems merely store rows; FlowVision actively optimizes and predicts.",
    "Smart Intent Routing: Categorizes queries (SYSTEM_TOPOLOGY, semantic_search) for instant context routing.",
    "Predictive Bottleneck Engine: Calculates transit momentum to predict workload spikes 4.2 hours in advance.",
    "Natural Language Querying: Enables non-technical clerks to query complex databases using plain English."
])

# Slide 4
add_content_slide(prs, "High-Performance Multi-Layer Architecture", [
    "Frontend & Backend: Vue.js, Nuxt.js, Tailwind CSS (Glassmorphic UI) + Node.js, Express.js.",
    "Database & Cloud: MySQL (Relational mesh), Supabase (Real-time sync), AWS S3 (QR metadata vaults).",
    "AI Pipeline: Groq / LLaMA (Ultra-low latency LLM inference) + Python background analytics.",
    "Flow Highlights: Smart QR Checkpoint scanning -> Cryptographic QR Hash -> Zero-Ambiguity Audit Log."
])

# Slide 5
add_content_slide(prs, "Privacy-First, Decentralized Governance", [
    "Federated Learning Network: AI models learn operational trends locally without centralizing citizen data.",
    "Strict Multi-Tenant Isolation: AI session memory (aiSession.ts) is hard-scoped to the user's org_id and user_id.",
    "Human-in-the-Loop & Immutable Ledger: AI provides predictions, but physical custody relies on immutable SQL logs and QR hashes."
])

# Slide 6
add_content_slide(prs, "Transforming Municipal Efficiency", [
    "Eradicating Lost Files: Eliminates document loss and citizen frustration across government offices.",
    "Accelerated Approval Pipelines: Faster processing for business licenses, permits, and clearances.",
    "Economic Scalability: Edge Node Cloud allows new municipal offices to join the routing mesh instantly."
])

# Slide 7
add_content_slide(prs, "Empowering Governments with Proactive Intelligence", [
    "Core Summary: Unifying isolated offices, couriers, and citizens into one transparent, measurable ecosystem.",
    "Closing Statement: Team Immortals — Transforming Municipal Workflows with FlowVision."
], footer="Ready for Q&A.")

prs.save("FlowVision_DASIG_Pitch.pptx")
